from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin, sec.bottom_margin = Inches(.68), Inches(.65)
sec.left_margin, sec.right_margin = Inches(.75), Inches(.75)
sec.header_distance, sec.footer_distance = Inches(.28), Inches(.28)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    style = doc.styles[name]
    style.font.name = 'Calibri'
    style.font.color.rgb = RGBColor(0,0,0)
    style.paragraph_format.widow_control = True
normal = doc.styles['Normal']
normal.font.size = Pt(10.5)
normal.paragraph_format.line_spacing = 1.10
normal.paragraph_format.space_after = Pt(6)
for name,size in [('Title',28),('Heading 1',18),('Heading 2',12),('Heading 3',11)]:
    style=doc.styles[name]
    style.font.size=Pt(size)
    style.font.bold = name != 'Title'
    style.paragraph_format.space_before=Pt(10 if name!='Title' else 0)
    style.paragraph_format.space_after=Pt(6)
    style.paragraph_format.keep_with_next=True
header=sec.header.paragraphs[0]
header.text='ACCESSPATH ULETH   /   PROJECT REPORT'
header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('Deep  |  Source snapshot b5073f7  |  ');r.font.size=Pt(8)
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='AccessPath ULeth Project Report'
doc.core_properties.subject='Technical implementation and future development'
doc.core_properties.author='Deep'
doc.core_properties.keywords='campus navigation accessibility JavaScript SVG MapLibre routing'

def inline(p,text):
    pos=0
    for m in re.finditer(r'\[([^\]]+)\]\(([^)]+)\)',text):
        p.add_run(text[pos:m.start()])
        link=OxmlElement('w:hyperlink')
        link.set(qn('r:id'),p.part.relate_to(m.group(2),RT.HYPERLINK,is_external=True))
        run=OxmlElement('w:r');props=OxmlElement('w:rPr')
        color=OxmlElement('w:color');color.set(qn('w:val'),'235A72');props.append(color)
        under=OxmlElement('w:u');under.set(qn('w:val'),'single');props.append(under)
        run.append(props);t=OxmlElement('w:t');t.text=m.group(1);run.append(t);link.append(run);p._p.append(link)
        pos=m.end()
    p.add_run(text[pos:])

def table(rows):
    t=doc.add_table(rows=0,cols=len(rows[0]))
    t.alignment=WD_TABLE_ALIGNMENT.CENTER
    t.autofit=False
    widths=[1.65,5.35] if len(rows[0])==2 else [2.1]*len(rows[0])
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'D9D9D9');borders.append(b)
    t._tbl.tblPr.append(borders)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        trPr=t.rows[-1]._tr.get_or_add_trPr()
        no_split=OxmlElement('w:cantSplit');trPr.append(no_split)
        if i==0:repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for j,(c,txt) in enumerate(zip(cells,row)):
            c.width=Inches(widths[j]);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=c._tc.get_or_add_tcPr();margins=OxmlElement('w:tcMar')
            for side,val in [('top',65),('bottom',65),('left',95),('right',95)]:
                el=OxmlElement('w:'+side);el.set(qn('w:w'),str(val));el.set(qn('w:type'),'dxa');margins.append(el)
            pr.append(margins)
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E6EEF2' if i==0 else ('F5F7F8' if i%2==0 else 'FFFFFF'));pr.append(shade)
            p=c.paragraphs[0];p.paragraph_format.space_after=Pt(0);p.paragraph_format.line_spacing=1.02
            inline(p,txt)
            for r in p.runs:r.font.size=Pt(9.5);r.font.bold=i==0
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(2);p.paragraph_format.space_before=Pt(0);p.paragraph_format.line_spacing=Pt(2);p.add_run().font.size=Pt(2)

lines=(ROOT/'AccessPath-ULeth-Project-Report.md').read_text(encoding='utf-8').splitlines()
i=0
pending_break=False
while i<len(lines):
    s=lines[i].strip()
    if not s:i+=1;continue
    if s=='<!-- PAGE -->':pending_break=True;i+=1;continue
    image_match=re.fullmatch(r'!\[([^\]]+)\]\(([^)]+)\)',s)
    if image_match:
        p=doc.add_paragraph()
        p.paragraph_format.keep_with_next=True
        p.paragraph_format.space_after=Pt(4)
        pic=p.add_run().add_picture(str(ROOT/image_match.group(2)),width=Inches(7))
        pic._inline.docPr.set('descr',image_match.group(1))
        i+=1
        continue
    if s.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            row=[v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.fullmatch(r':?-+:?',v) for v in row):rows.append(row)
            i+=1
        table(rows);continue
    if s.startswith('# '):inline(doc.add_paragraph(style='Title'),s[2:])
    elif s.startswith('## '):inline(doc.add_paragraph(style='Heading 1'),s[3:])
    elif s.startswith('### '):inline(doc.add_paragraph(style='Heading 2'),s[4:])
    elif s.startswith('- '):
        p=doc.add_paragraph(style='List Bullet');p.paragraph_format.space_after=Pt(5);inline(p,s[2:])
    else:inline(doc.add_paragraph(),s)
    if pending_break:
        doc.paragraphs[-1].paragraph_format.page_break_before=True
        pending_break=False
    i+=1
for border in list(doc.styles.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
for border in list(doc.element.iter(qn('w:pBdr'))):
    border.getparent().remove(border)
doc.save(ROOT/'AccessPath-ULeth-Project-Report.docx')
print('Saved project report DOCX')


